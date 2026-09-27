"""Fill a PowerPoint template: replace text and pictures of named shapes on given slides, keep everything else.

Usage: python build_deck.py spec.json
spec = {
  "template": "in.pptx", "output": "out.pptx",
  "slides": [{"index": 31,                                  # 1-based slide number
              "texts": {"<shape name>": "new text" | ["para 1", null, ...]},  # keeps each paragraph's first-run
                                                            # formatting; null keeps that paragraph as is
              "images": {"<shape name>": {"path": "a.png", "box": [left, top, width, height] | null}}}]
}
Pictures are re-embedded as new image parts (images shared with other slides stay untouched) and fitted
inside the original box (or `box`, in inches) without stretching.
"""
import json
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches
from PIL import Image


def find_shape(shapes, name):
    for shape in shapes:
        if shape.name == name:
            return shape
        if shape.shape_type == 6:  # group: search inside
            found = find_shape(shape.shapes, name)
            if found is not None:
                return found
    return None


def set_paragraph(paragraph, value):
    value = " ".join(str(value).split())
    runs = paragraph.runs
    if not runs:
        paragraph.text = value
        return
    runs[0].text = value
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    for br in paragraph._p.findall(qn("a:br")):
        paragraph._p.remove(br)


def set_text(shape, value):
    """A string replaces the whole box (one paragraph); a list sets paragraph by paragraph."""
    if not shape.has_text_frame:
        raise ValueError(f"shape {shape.name!r} has no text")
    paragraphs = shape.text_frame.paragraphs
    values = value if isinstance(value, list) else [value]
    for extra in paragraphs[len(values):]:
        extra._p.getparent().remove(extra._p)
    for paragraph, text in zip(paragraphs, values):
        if text is not None:
            set_paragraph(paragraph, text)


def replace_picture(slide, shape, path, box):
    if shape.shape_type != 13:
        raise ValueError(f"shape {shape.name!r} is not a picture")
    _, rid = slide.part.get_or_add_image_part(path)
    blip_fill = shape._element.blipFill
    blip_fill.blip.set(qn("r:embed"), rid)
    for crop in blip_fill.findall(qn("a:srcRect")):  # template crops don't apply to the new image
        blip_fill.remove(crop)

    left, top, width, height = (
        (Inches(box[0]), Inches(box[1]), Inches(box[2]), Inches(box[3])) if box else (shape.left, shape.top, shape.width, shape.height)
    )
    with Image.open(path) as img:
        img_w, img_h = img.size
    scale = min(width / img_w, height / img_h)
    new_w, new_h = int(img_w * scale), int(img_h * scale)
    shape.left = int(left + (width - new_w) / 2)
    shape.top = int(top + (height - new_h) / 2)
    shape.width = new_w
    shape.height = new_h


def main(spec_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation(spec["template"])
    total = len(prs.slides)
    for item in spec["slides"]:
        index = item["index"]
        if index < 1 or index > total:
            raise ValueError(f"template has {total} slides, cannot edit slide {index}")
        slide = prs.slides[index - 1]
        for name, value in item.get("texts", {}).items():
            shape = find_shape(slide.shapes, name)
            if shape is None:
                raise ValueError(f"slide {index}: shape {name!r} not found — is this the right template?")
            set_text(shape, value)
        for name, image in item.get("images", {}).items():
            shape = find_shape(slide.shapes, name)
            if shape is None:
                raise ValueError(f"slide {index}: picture {name!r} not found — is this the right template?")
            replace_picture(slide, shape, image["path"], image.get("box"))
    prs.save(spec["output"])


if __name__ == "__main__":
    try:
        main(sys.argv[1])
    except Exception as exc:  # message goes to the Node side, which shows it to the user
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        sys.exit(1)
